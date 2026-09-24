import io
import json
import sys
import uuid
from pathlib import Path

# Ensure local forge package takes precedence
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest

from forge.converter import (
    convert_and_upload,
    convert_bytes,
    convert_bytes_with_metadata,
    convert_directory,
    convert_directory_and_upload,
    convert_file,
    convert_to_cleaned,
    convert_to_parsed,
    ingest_converted,
)
from forge.ingest import Cleaned, ParseError, parse, parse_json, parse_with_converter
from forge.models import Bucket, Document
from forge.ocr import (
    get_ocr_engine_name,
    is_ocr_available,
    ocr_document,
    ocr_image,
    ocr_pdf,
)

FIXTURES = Path(__file__).parent / "fixtures"
DESKTOP_INPUTS = Path(r"C:\Users\PSK\Desktop\pdf-docx-to-json\input_files")


def test_convert_file_plain_text(tmp_path: Path):
    txt_file = tmp_path / "test.txt"
    txt_file.write_text("Hello, Docling conversion!", encoding="utf-8")

    record = convert_file(txt_file)
    assert record["filename"] == "test.txt"
    assert "Hello, Docling conversion!" in record["text"]


def test_convert_file_markdown(tmp_path: Path):
    md_file = tmp_path / "notes.md"
    md_file.write_text("# Title\n\n- item 1\n- item 2", encoding="utf-8")

    record = convert_file(md_file)
    assert record["filename"] == "notes.md"
    assert "# Title" in record["text"]


def test_convert_to_parsed(tmp_path: Path):
    txt_file = tmp_path / "parsed_test.txt"
    txt_file.write_text("Parsed content here", encoding="utf-8")

    parsed = convert_to_parsed(txt_file)
    assert parsed.filename == "parsed_test.txt"
    assert parsed.text == "Parsed content here"


def test_convert_directory(tmp_path: Path):
    in_dir = tmp_path / "inputs"
    in_dir.mkdir()
    (in_dir / "a.txt").write_text("Alpha text", encoding="utf-8")
    (in_dir / "b.md").write_text("Beta text", encoding="utf-8")
    (in_dir / "c.csv").write_text("Col1,Col2", encoding="utf-8")

    out_json = tmp_path / "output.json"
    records = convert_directory(in_dir, output_file=out_json)

    assert len(records) == 2
    filenames = {r["filename"] for r in records}
    assert filenames == {"a.txt", "b.md"}

    assert out_json.is_file()
    saved_records = json.loads(out_json.read_text(encoding="utf-8"))
    assert len(saved_records) == 2


def test_convert_bytes_txt():
    text = convert_bytes("sample.txt", b"In-memory raw text bytes")
    assert text == "In-memory raw text bytes"


def test_convert_file_unsupported(tmp_path: Path):
    bad = tmp_path / "data.csv"
    bad.write_text("1,2,3", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported file format"):
        convert_file(bad)


def test_convert_file_missing():
    with pytest.raises(FileNotFoundError):
        convert_file("non_existent_file.docx")


def test_parse_json_from_converter():
    records = [
        {"filename": "doc1.pdf", "text": "Extracted text from doc1"},
        {"filename": "doc2.docx", "text": "Extracted text from doc2"},
    ]
    raw = json.dumps(records).encode("utf-8")
    parsed_docs = parse_json(raw, "output.json")

    assert len(parsed_docs) == 2
    assert parsed_docs[0].filename == "doc1.pdf"
    assert parsed_docs[0].text == "Extracted text from doc1"
    assert parsed_docs[1].filename == "doc2.docx"
    assert parsed_docs[1].text == "Extracted text from doc2"


def test_parse_json_via_ingest_parse():
    records = [{"filename": "sample.pdf", "text": "Sample text content"}]
    raw = json.dumps(records).encode("utf-8")

    docs = parse("output.json", raw)
    assert len(docs) == 1
    assert docs[0].filename == "sample.pdf"
    assert docs[0].text == "Sample text content"


def test_parse_json_errors():
    with pytest.raises(ParseError, match="could not be decoded"):
        parse_json(b"invalid-json{", "broken.json")

    with pytest.raises(ParseError, match="No usable records"):
        parse_json(json.dumps([{"no_text": 123}]).encode("utf-8"), "empty.json")


def test_parse_with_converter_txt():
    docs = parse_with_converter("sample.txt", b"Direct converter parse text")
    assert len(docs) == 1
    assert docs[0].text == "Direct converter parse text"
    assert docs[0].filename == "sample.txt"


def test_convert_image_file_with_ocr(tmp_path: Path):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (300, 100), color="white")
    d = ImageDraw.Draw(img)
    d.text((10, 30), "Docling Image OCR", fill="black")
    img_path = tmp_path / "test_ocr.png"
    img.save(img_path)

    record = convert_file(img_path)
    assert record["filename"] == "test_ocr.png"
    assert any(term in record["text"] for term in ("Docling", "Image", "OCR"))


def test_convert_bytes_image_with_ocr():
    import io
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (300, 100), color="white")
    d = ImageDraw.Draw(img)
    d.text((10, 30), "Docling In-Memory OCR", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    text = convert_bytes("test_mem.png", buf.getvalue())
    assert any(term in text for term in ("Docling", "Memory", "OCR"))


@pytest.mark.slow
def test_docling_docx_conversion_if_sample_exists():
    docx_sample = DESKTOP_INPUTS / "sample.docx"
    if not docx_sample.is_file():
        pytest.skip(f"Sample docx not found at {docx_sample}")

    record = convert_file(docx_sample)
    assert record["filename"] == "sample.docx"
    assert len(record["text"]) > 0
    assert "Sample Word Document" in record["text"]


def test_convert_to_cleaned(tmp_path: Path):
    doc_path = tmp_path / "cleaned_sample.txt"
    doc_path.write_text(
        "A thoroughly ordinary paragraph of English prose that clears all cleaning thresholds. "
        * 10,
        encoding="utf-8",
    )
    cleaned = convert_to_cleaned(doc_path)
    assert isinstance(cleaned, Cleaned)
    assert cleaned.char_count > 0
    assert cleaned.word_count > 0
    assert cleaned.content_hash
    assert cleaned.quality["flags"] == []
    assert cleaned.quality["alpha_ratio"] > 0.5


def test_convert_file_with_clean_text(tmp_path: Path):
    doc_path = tmp_path / "report.md"
    doc_path.write_text("# Report Title\n\n\n\n- point 1\n- point 2   \n", encoding="utf-8")
    record = convert_file(doc_path, clean_text=True)

    assert record["filename"] == "report.md"
    assert "content_hash" in record
    assert "quality" in record
    assert "char_count" in record
    assert "word_count" in record
    assert record["text"] == "# Report Title\n\n- point 1\n- point 2"


def test_convert_directory_with_clean_text(tmp_path: Path):
    in_dir = tmp_path / "batch_in"
    in_dir.mkdir()
    (in_dir / "1.txt").write_text("First file line.   \n\n\n\nSecond line.", encoding="utf-8")
    (in_dir / "2.md").write_text("# Second File\n\nDescription text.", encoding="utf-8")

    out_json = tmp_path / "batch_out.json"
    records = convert_directory(in_dir, output_file=out_json, clean_text=True)

    assert len(records) == 2
    for r in records:
        assert "content_hash" in r
        assert "quality" in r
        assert "char_count" in r


def test_convert_and_upload_to_bucket(tmp_path: Path, db):
    bucket = Bucket(name="test_upload_bucket")
    db.add(bucket)
    db.commit()

    doc_file = tmp_path / "upload_sample.txt"
    doc_file.write_text(
        "Ingested converted document content that clears standard thresholds. " * 8,
        encoding="utf-8",
    )

    doc = convert_and_upload(
        doc_file,
        bucket=bucket.id,
        session=db,
        source_note="test run",
    )

    assert doc.status == "parsed"
    assert doc.filename == "upload_sample.txt"
    assert doc.source_format == "txt"
    assert doc.source_note == "test run"
    assert doc.char_count > 0
    assert doc.text_uri is not None


def test_convert_and_upload_duplicate_detection(tmp_path: Path, db):
    bucket = Bucket(name="test_dup_bucket")
    db.add(bucket)
    db.commit()

    text_content = "Identical text content for duplicate testing. " * 10
    f1 = tmp_path / "original.txt"
    f1.write_text(text_content, encoding="utf-8")
    f2 = tmp_path / "copy.txt"
    f2.write_text(text_content, encoding="utf-8")

    doc1 = convert_and_upload(f1, bucket=bucket.id, session=db)
    assert doc1.status == "parsed"

    doc2 = convert_and_upload(f2, bucket=bucket.id, session=db)
    assert doc2.status == "failed"
    assert "Duplicate of original.txt" in doc2.error


def test_convert_directory_and_upload(tmp_path: Path, db):
    bucket = Bucket(name="test_batch_bucket")
    db.add(bucket)
    db.commit()

    in_dir = tmp_path / "docs_dir"
    in_dir.mkdir()
    (in_dir / "doc_a.txt").write_text("Alpha unique text line. " * 10, encoding="utf-8")
    (in_dir / "doc_b.txt").write_text("Beta unique text line. " * 10, encoding="utf-8")

    docs = convert_directory_and_upload(in_dir, bucket=bucket.name, session=db)
    assert len(docs) == 2
    assert all(d.status == "parsed" for d in docs)


def test_convert_docx_paragraphs_and_tables(tmp_path: Path):
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("Docx Section Header", level=2)
    doc.add_paragraph("Paragraph before table.")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Product"
    tbl.cell(0, 1).text = "Price"
    tbl.cell(1, 0).text = "Widget"
    tbl.cell(1, 1).text = "$9.99"
    doc.add_paragraph("Paragraph after table.")

    docx_path = tmp_path / "sample_table.docx"
    doc.save(str(docx_path))

    record = convert_file(docx_path, clean_text=True)
    assert record["filename"] == "sample_table.docx"
    assert "## Docx Section Header" in record["text"]
    assert "Paragraph before table." in record["text"]
    assert "| Product | Price |" in record["text"]
    assert "| Widget | $9.99 |" in record["text"]
    assert "Paragraph after table." in record["text"]


def test_convert_bytes_docx_table():
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("In-Memory Docx", level=1)
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "A"
    tbl.cell(0, 1).text = "B"
    tbl.cell(1, 0).text = "1"
    tbl.cell(1, 1).text = "2"
    buf = io.BytesIO()
    doc.save(buf)

    text = convert_bytes("memory.docx", buf.getvalue())
    assert "# In-Memory Docx" in text
    assert "| A | B |" in text
    assert "| 1 | 2 |" in text


def test_ocr_availability():
    # If test suite runs, check engine name representation
    name = get_ocr_engine_name()
    assert isinstance(name, str)
    assert len(name) > 0


def test_ocr_jpg_and_jpeg_images(tmp_path: Path):
    from PIL import Image, ImageDraw

    # Test JPG
    img_jpg = Image.new("RGB", (320, 100), color="white")
    d1 = ImageDraw.Draw(img_jpg)
    d1.text((10, 30), "Invoice JPG Scan 2026", fill="black")
    jpg_path = tmp_path / "receipt.jpg"
    img_jpg.save(str(jpg_path), format="JPEG")

    text_jpg = ocr_image(jpg_path)
    assert any(term in text_jpg for term in ("Invoice", "Scan", "2026"))

    # Test JPEG via convert_file
    img_jpeg = Image.new("RGB", (320, 100), color="white")
    d2 = ImageDraw.Draw(img_jpeg)
    d2.text((10, 30), "Warranty JPEG 2026", fill="black")
    jpeg_path = tmp_path / "warranty.jpeg"
    img_jpeg.save(str(jpeg_path), format="JPEG")

    record = convert_file(jpeg_path, clean_text=True)
    assert record["filename"] == "warranty.jpeg"
    assert any(term in record["text"] for term in ("Warranty", "2026"))


def test_ocr_scanned_pdf(tmp_path: Path):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (360, 120), color="white")
    d = ImageDraw.Draw(img)
    d.text((10, 30), "Scanned Legal Deed 2026", fill="black")
    pdf_path = tmp_path / "scanned_deed.pdf"
    img.save(str(pdf_path), format="PDF")

    text = ocr_pdf(pdf_path)
    assert any(term in text for term in ("Scanned", "Legal", "Deed", "2026"))


def test_ocr_document_dispatcher(tmp_path: Path):
    from PIL import Image, ImageDraw

    # Test PNG through unified dispatcher
    img = Image.new("RGB", (300, 100), color="white")
    d = ImageDraw.Draw(img)
    d.text((10, 30), "Dispatcher PNG OCR", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    text = ocr_document(buf.getvalue(), filename="test.png")
    assert any(term in text for term in ("Dispatcher", "PNG", "OCR"))


def test_convert_file_plain_text_metadata(tmp_path: Path):
    txt_file = tmp_path / "metadata_test.txt"
    txt_file.write_text("Hello, extraction metadata!\nSecond line of text.", encoding="utf-8")

    record = convert_file(txt_file, include_metadata=True)
    assert "metadata" in record
    meta = record["metadata"]
    assert meta["source_format"] == "txt"
    assert meta["engine"] == "native"
    assert meta["ocr_applied"] is False
    assert meta["char_count"] == len(record["text"])
    assert meta["word_count"] == len(record["text"].split())
    assert meta["line_count"] == 2


def test_convert_file_disable_metadata(tmp_path: Path):
    txt_file = tmp_path / "no_metadata.txt"
    txt_file.write_text("No metadata requested.", encoding="utf-8")

    record = convert_file(txt_file, include_metadata=False)
    assert "metadata" not in record
    assert record["filename"] == "no_metadata.txt"
    assert record["text"] == "No metadata requested."


def test_convert_file_docx_fixture_metadata():
    docx_fixture = FIXTURES / "sample.docx"
    assert docx_fixture.is_file()

    record = convert_file(docx_fixture, include_metadata=True)
    assert record["filename"] == "sample.docx"
    assert "# Forge Architecture Overview" in record["text"]
    assert "## Core Components" in record["text"]
    assert "| Component | Throughput | Status |" in record["text"]

    meta = record["metadata"]
    assert meta["source_format"] == "docx"
    assert meta["engine"] == "openxml"
    assert meta["ocr_applied"] is False
    assert meta["table_count"] == 1
    assert meta["heading_count"] >= 3
    assert meta["paragraph_count"] >= 5


def test_convert_file_pdf_fixture_metadata():
    pdf_fixture = FIXTURES / "sample.pdf"
    assert pdf_fixture.is_file()

    record = convert_file(pdf_fixture, include_metadata=True)
    assert record["filename"] == "sample.pdf"
    assert "Forge Document Ingestion" in record["text"]

    meta = record["metadata"]
    assert meta["source_format"] == "pdf"
    assert meta["engine"] == "pypdf"
    assert meta["ocr_applied"] is False
    assert meta["page_count"] >= 1


def test_convert_file_image_png_fixture_metadata():
    png_fixture = FIXTURES / "sample.png"
    assert png_fixture.is_file()

    record = convert_file(png_fixture, include_metadata=True)
    assert record["filename"] == "sample.png"
    assert "metadata" in record
    meta = record["metadata"]
    assert meta["source_format"] == "png"
    assert meta["ocr_applied"] is True
    assert meta["width"] == 100
    assert meta["height"] == 40


def test_convert_file_image_jpg_fixture_metadata():
    jpg_fixture = FIXTURES / "sample.jpg"
    assert jpg_fixture.is_file()

    record = convert_file(jpg_fixture, include_metadata=True)
    assert record["filename"] == "sample.jpg"
    assert "metadata" in record
    meta = record["metadata"]
    assert meta["source_format"] in ("jpg", "jpeg")
    assert meta["ocr_applied"] is True
    assert meta["width"] == 1
    assert meta["height"] == 1


def test_convert_bytes_with_metadata_docx():
    docx_fixture = FIXTURES / "sample.docx"
    raw = docx_fixture.read_bytes()

    text, meta = convert_bytes_with_metadata("sample.docx", raw)
    assert "# Forge Architecture Overview" in text
    assert "| Component | Throughput | Status |" in text
    assert meta["source_format"] == "docx"
    assert meta["engine"] == "openxml"
    assert meta["table_count"] == 1


def test_convert_to_parsed_preserves_metadata():
    docx_fixture = FIXTURES / "sample.docx"
    parsed = convert_to_parsed(docx_fixture)

    assert parsed.filename == "sample.docx"
    assert "# Forge Architecture Overview" in parsed.text
    assert parsed.metadata["source_format"] == "docx"
    assert parsed.metadata["engine"] == "openxml"
    assert parsed.metadata["table_count"] == 1


def test_convert_directory_with_fixtures(tmp_path: Path):
    in_dir = tmp_path / "fixture_inputs"
    in_dir.mkdir()

    import shutil

    shutil.copy(FIXTURES / "sample.docx", in_dir / "sample.docx")
    shutil.copy(FIXTURES / "sample.pdf", in_dir / "sample.pdf")
    shutil.copy(FIXTURES / "plain.txt", in_dir / "plain.txt")

    out_json = tmp_path / "fixtures_converted.json"
    records = convert_directory(
        in_dir, output_file=out_json, clean_text=True, include_metadata=True
    )

    assert len(records) == 3
    filenames = {r["filename"] for r in records}
    assert filenames == {"sample.docx", "sample.pdf", "plain.txt"}

    for r in records:
        assert "metadata" in r
        assert "quality" in r
        assert "content_hash" in r
        assert r["metadata"]["char_count"] > 0

    assert out_json.is_file()
    saved = json.loads(out_json.read_text(encoding="utf-8"))
    assert len(saved) == 3


def test_ingest_converted_stores_extraction_metadata(db):
    docx_fixture = FIXTURES / "sample.docx"
    record = convert_file(docx_fixture, include_metadata=True)

    bucket = Bucket(name=f"metadata-test-{uuid.uuid4().hex[:6]}", description="test metadata bucket")
    db.add(bucket)
    db.commit()

    doc = ingest_converted(
        session=db,
        bucket_id=bucket.id,
        filename=record["filename"],
        text=record["text"],
        metadata=record.get("metadata"),
    )
    db.commit()

    assert doc.status == "parsed"
    assert "extraction" in doc.quality
    ext_meta = doc.quality["extraction"]
    assert ext_meta["source_format"] == "docx"
    assert ext_meta["engine"] == "openxml"
    assert ext_meta["table_count"] == 1




