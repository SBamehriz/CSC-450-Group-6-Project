import gzip
from pathlib import Path

import pytest

from forge.ingest import (
    ParseError,
    check_supported,
    clean,
    drop_repeated_lines,
    flag,
    looks_mis_decoded,
    measure,
    normalise,
    parse,
    split_extension,
)
from tests.conftest import ARTICLE_PDF, NOTES_MD, SCANNED_PDF

FIXTURES = Path(__file__).parent / "fixtures"


GENERATED = {"article.pdf": ARTICLE_PDF, "scanned.pdf": SCANNED_PDF, "notes.md": NOTES_MD}


def read(name: str) -> bytes:
    if name in GENERATED:
        return GENERATED[name]
    return (FIXTURES / name).read_bytes()


def golden(name: str, produced: str) -> None:
    # delete the file to accept a change
    expected = FIXTURES / f"{name}.expected.txt"
    if not expected.exists():
        expected.write_text(produced, encoding="utf-8")
        pytest.skip(f"wrote a new snapshot {expected.name} — check it and re-run")
    assert produced == expected.read_text(encoding="utf-8")


@pytest.mark.parametrize(
    "filename,expected",
    [
        ("a.txt", ("txt", False)),
        ("a.TXT", ("txt", False)),
        ("notes.md", ("md", False)),
        ("page.HTM", ("htm", False)),
        ("dump.jsonl.gz", ("jsonl", True)),
        ("weird.name.with.dots.txt", ("txt", False)),
    ],
)
def test_split_extension(filename, expected):
    assert split_extension(filename) == expected


def test_htm_is_html():
    assert check_supported("page.htm") == "html"


def test_unsupported_extension_names_what_we_take():
    with pytest.raises(ParseError) as caught:
        check_supported("archive.zip")
    message = str(caught.value)
    assert ".zip" in message
    for good in (".txt", ".md", ".html", ".pdf", ".docx", ".png", ".jsonl"):
        assert good in message


def test_uploading_an_unsupported_file_fails_the_same_way():
    with pytest.raises(ParseError):
        parse("archive.zip", b"fake zip bytes")


def test_uploading_corrupted_docx_fails():
    with pytest.raises(ParseError):
        parse("thesis.docx", read("unsupported.docx"))


def test_docx_parse_success():
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("Docx Ingestion Header", level=1)
    doc.add_paragraph("This is body text in a docx file.")
    buf = io.BytesIO()
    doc.save(buf)

    [document] = parse("notes.docx", buf.getvalue())
    assert "Docx Ingestion Header" in document.text
    assert "body text in a docx file" in document.text
    assert document.filename == "notes.docx"


def test_image_parse_with_ocr():
    import io
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (320, 100), color="white")
    d = ImageDraw.Draw(img)
    d.text((10, 30), "Invoice Receipt 2026", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    [document] = parse("invoice.png", buf.getvalue())
    assert any(w in document.text for w in ("Invoice", "Receipt", "2026"))
    assert document.filename == "invoice.png"


def test_scanned_pdf_ocr_success():
    import io
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (320, 100), color="white")
    d = ImageDraw.Draw(img)
    d.text((10, 30), "Scanned Contract Page", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PDF")

    [document] = parse("contract_scan.pdf", buf.getvalue())
    assert any(w in document.text for w in ("Scanned", "Contract", "Page"))
    assert document.filename == "contract_scan.pdf"


def test_plain_text():
    [document] = parse("plain.txt", read("plain.txt"))
    golden("plain.txt", clean(document.text).text)


def test_markdown_is_kept_as_is():
    [document] = parse("notes.md", read("notes.md"))
    cleaned = clean(document.text)
    assert "**as-is**" in cleaned.text
    assert "## What we do change" in cleaned.text
    golden("notes.md", cleaned.text)


def test_html_loses_its_boilerplate():
    [document] = parse("article.html", read("article.html"))
    cleaned = clean(document.text)
    assert "Copyright 2026" not in cleaned.text
    assert "Privacy" not in cleaned.text
    assert "Training compute scales" in cleaned.text
    golden("article.html", cleaned.text)


def test_pdf_text_layer():
    [document] = parse("article.pdf", read("article.pdf"))
    cleaned = clean(document.text)
    assert "quick brown fox" in cleaned.text
    golden("article.pdf", cleaned.text)


def test_gzip_unwraps_to_the_same_thing():
    raw = b"Gzipped text works exactly like plain text once unpacked.\n" * 20
    [from_gz] = parse("compressed.txt.gz", read("compressed.txt.gz"))
    [from_plain] = parse("compressed.txt", raw)
    assert clean(from_gz.text).content_hash == clean(from_plain.text).content_hash
    assert from_gz.filename == "compressed.txt"


def test_jsonl_fans_out_one_document_per_line():
    documents = parse("records.jsonl", read("records.jsonl"))
    # 6 lines, 3 usable
    assert len(documents) == 3
    assert [d.filename for d in documents] == [
        "records.jsonl#L1",
        "records.jsonl#L2",
        "records.jsonl#L6",
    ]
    golden("records.jsonl", clean(documents[0].text).text)


def test_scanned_pdf_is_refused_in_plain_words():
    with pytest.raises(ParseError) as caught:
        parse("scanned.pdf", read("scanned.pdf"))
    message = str(caught.value)
    assert "scanned" in message.lower()
    assert "OCR" in message


def test_jsonl_with_no_usable_lines_says_why():
    with pytest.raises(ParseError) as caught:
        parse("broken.jsonl", read("broken.jsonl"))
    assert "No usable lines" in str(caught.value)


def test_empty_text_file_is_refused():
    with pytest.raises(ParseError, match="empty"):
        parse("blank.txt", b"   \n\n  ")


def test_html_with_no_article_is_refused():
    with pytest.raises(ParseError, match="No article content"):
        parse("nav.html", b"<html><body><nav><a href='/'>home</a></nav></body></html>")


def test_corrupt_gzip_is_refused():
    with pytest.raises(ParseError, match="could not be unpacked"):
        parse("bad.txt.gz", b"this is not gzip data at all")


def test_corrupt_pdf_is_refused():
    with pytest.raises(ParseError, match="could not be read"):
        parse("bad.pdf", b"%PDF-1.4 and then nothing valid")


def test_latin1_file_still_reads():
    [document] = parse("latin1.txt", read("latin1.txt"))
    assert "café" in document.text
    assert looks_mis_decoded(read("latin1.txt")) is True


def test_utf8_is_not_flagged_as_suspect():
    assert looks_mis_decoded("café".encode()) is False


def test_encoding_suspect_becomes_a_flag():
    cleaned = clean("some text " * 100, encoding_suspect=True)
    assert "encoding_suspect" in cleaned.quality["flags"]


def test_normalise_collapses_blank_runs_and_trailing_space():
    assert normalise("a   \n\n\n\n\nb  ") == "a\n\nb"


def test_normalise_unifies_newlines():
    assert normalise("a\r\nb\rc") == "a\nb\nc"


def test_normalise_strips_control_characters_but_keeps_tabs():
    assert normalise("a\x00b\tc\x07d") == "ab\tcd"


def test_normalise_is_unicode_nfc():
    # combining acute becomes one é
    assert normalise("café") == "café"


def test_repeated_short_lines_are_dropped():
    text = "\n".join(["Home"] * 6 + ["Real sentence here."] + ["Home"] * 2)
    assert "Home" not in drop_repeated_lines(text)
    assert "Real sentence here." in drop_repeated_lines(text)


def test_repeated_long_lines_are_kept():
    line = "This is a genuinely long repeated sentence that is part of the text. " * 2
    text = "\n".join([line] * 6)
    assert line.strip() in drop_repeated_lines(text)


def test_a_line_repeated_only_a_few_times_survives():
    text = "\n".join(["Chorus"] * 4 + ["verse"])
    assert "Chorus" in drop_repeated_lines(text)


def test_measure_counts_what_it_says():
    stats = measure("abc 123\nxyz")
    assert stats["alpha_ratio"] == round(6 / 11, 3)
    assert stats["digit_ratio"] == round(3 / 11, 3)
    assert stats["line_count"] == 2


def test_flags_fire_on_their_thresholds():
    assert "too_short" in flag("tiny", measure("tiny"))
    numbers = "1234567890 " * 200
    assert "low_alpha" in flag(numbers, measure(numbers))
    long_line = "word " * 500
    assert "long_lines" in flag(long_line, measure(long_line))


def test_a_normal_document_has_no_flags():
    text = "This is an ordinary paragraph of English prose. " * 30
    assert flag(text, measure(text)) == []


def test_same_input_gives_the_same_hash_every_time():
    text = "Determinism is the whole point of this pipeline.\n\n\nReally."
    assert clean(text).content_hash == clean(text).content_hash


def test_cleaning_makes_two_sources_collide_when_only_whitespace_differs():
    a = clean("Hello world.\n\n\n\nSecond.")
    b = clean("Hello world.\n\nSecond.   ")
    assert a.content_hash == b.content_hash


def test_different_text_gives_a_different_hash():
    assert clean("one").content_hash != clean("two").content_hash


def test_cleaning_none_keeps_the_raw_text():
    messy = "a   \n\n\n\n\nb  "
    assert clean(messy, mode="none").text == messy
    assert clean(messy, mode="none").content_hash != clean(messy).content_hash


def test_unknown_cleaning_mode_is_an_error():
    with pytest.raises(ValueError, match="unknown cleaning mode"):
        clean("x", mode="aggressive")


def test_stats_are_computed_even_without_cleaning():
    cleaned = clean("short", mode="none")
    assert cleaned.quality["alpha_ratio"] > 0
    assert "too_short" in cleaned.quality["flags"]


def test_gzip_round_trip_is_byte_identical():
    original = b"round trip me\n" * 50
    [document] = parse("x.txt.gz", gzip.compress(original))
    assert document.text == original.decode()
