from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from sqlalchemy.exc import OperationalError

from forge.db import get_session
from forge.models import Bucket


def test_empty_overview(client):
    stats = client.get("/api/overview/stats")
    assert stats.status_code == 200
    assert all(value == 0 for value in stats.json().values())
    assert client.get("/api/overview/recent-uploads").json() == {"items": [], "total": 0}


def test_stats_count_current_status_and_only_parsed_text(client, db, make_document, make_dataset):
    bucket = Bucket(name="Library")
    db.add(bucket)
    db.commit()
    parsed = make_document(bucket.id, chars=120)
    parsed.quality = {"flags": ["too_short"], "extraction": {"engine": "openxml"}}
    make_document(bucket.id, chars=80)
    make_document(bucket.id, chars=500, status="failed")
    make_document(bucket.id, chars=900, status="rejected")
    make_document(bucket.id, chars=700, status="pending")
    make_dataset(status="ready")
    make_dataset(status="draft")
    db.commit()
    assert client.get("/api/overview/stats").json() == {
        "buckets": 1,
        "documents": 5,
        "parsed_documents": 2,
        "failed_documents": 1,
        "rejected_documents": 1,
        "pending_documents": 1,
        "flagged_documents": 1,
        "chars": 200,
        "words": 40,
        "est_tokens": 50,
        "snapshots": 2,
        "ready_snapshots": 1,
    }


def test_recent_uploads_have_stable_pages_and_extraction_details(client, db, make_document):
    bucket = Bucket(name="Sources")
    db.add(bucket)
    db.commit()
    documents = [make_document(bucket.id) for _ in range(5)]
    same_time = datetime(2026, 10, 2, tzinfo=UTC)
    for document in documents:
        document.created_at = same_time
        document.quality = {"extraction": {"engine": "openxml", "ocr_applied": False}}
    db.commit()
    first = client.get("/api/overview/recent-uploads?limit=2").json()
    second = client.get("/api/overview/recent-uploads?limit=3&offset=2").json()
    ids = [item["id"] for item in first["items"] + second["items"]]
    assert ids == [str(doc.id) for doc in sorted(documents, key=lambda d: d.id, reverse=True)]
    assert first["total"] == second["total"] == 5
    assert all(item["bucket_name"] == "Sources" for item in first["items"])
    assert first["items"][0]["quality"]["extraction"]["ocr_applied"] is False
    assert client.get("/api/overview/recent-uploads?offset=20").json() == {"items": [], "total": 5}
    doc_first = client.get(f"/api/documents?bucket_id={bucket.id}&limit=2").json()
    doc_next = client.get(f"/api/documents?bucket_id={bucket.id}&limit=3&offset=2").json()
    assert [item["id"] for item in doc_first["items"] + doc_next["items"]] == ids


@pytest.mark.parametrize("query", ["limit=0", "limit=51", "offset=-1", "limit=abc"])
def test_recent_upload_validation(client, query):
    response = client.get(f"/api/overview/recent-uploads?{query}")
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_request"


@pytest.mark.parametrize("path", ["/api/overview/stats", "/api/overview/recent-uploads"])
def test_database_failure_is_an_api_error(client, path):
    class UnavailableSession:
        def execute(self, *args, **kwargs):
            raise OperationalError("SELECT", {}, Exception("offline"))

        scalar = execute

    client.app.dependency_overrides[get_session] = lambda: UnavailableSession()
    try:
        response = client.get(path)
        assert response.status_code == 503
        assert response.json()["error"]["code"] == "unavailable"
        assert "offline" not in response.text
    finally:
        client.app.dependency_overrides.clear()


def test_upload_review_snapshot_and_activity_refresh(client):
    bucket = client.post("/api/buckets", json={"name": "Sprint 5"}).json()
    docx = DocxDocument()
    docx.add_heading("Project notes", 1)
    docx.add_paragraph("A real source document for reviewing extraction metadata and text.")
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Datasets"
    output = BytesIO()
    docx.save(output)
    response = client.post(
        f"/api/buckets/{bucket['id']}/documents",
        files=[
            ("files", ("notes.docx", output.getvalue())),
            ("files", ("second.txt", b"A different document for the train and validation split.")),
            ("files", ("bad.exe", b"unsupported")),
        ],
        data={"source_note": "Class demo"},
    )
    assert response.status_code == 201
    assert response.json()["parsed"] == 2
    assert response.json()["failed"] == 1
    recent = client.get("/api/overview/recent-uploads").json()["items"]
    document = next(item for item in recent if item["filename"] == "notes.docx")
    assert document["quality"]["extraction"]["table_count"] == 1
    assert document["quality"]["extraction"]["engine"] == "openxml"
    assert document["source_note"] == "Class demo"
    detail = client.get(f"/api/documents/{document['id']}").json()
    assert detail["id"] == document["id"]
    text = client.get(f"/api/documents/{document['id']}/text?length=10").json()
    next_text = client.get(f"/api/documents/{document['id']}/text?offset=10&length=10").json()
    assert text["length"] == next_text["length"] == 10
    assert text["text"] != next_text["text"]
    snapshot = client.post(
        "/api/datasets",
        json={
            "name": "Sprint 5 snapshot",
            "bucket_id": bucket["id"],
            "validation_fraction": 0.5,
            "seed": 42,
        },
    )
    assert snapshot.status_code == 201
    assert client.get("/api/overview/stats").json()["ready_snapshots"] == 1
    assert client.get(f"/api/datasets/{snapshot.json()['id']}/download").status_code == 200
    client.post(f"/api/documents/{document['id']}/reject", json={"reason": "Review example"})
    assert client.get("/api/overview/stats").json()["parsed_documents"] == 1
    client.delete(f"/api/documents/{document['id']}")
    assert len(client.get("/api/overview/recent-uploads").json()["items"]) == 2
    assert client.get(f"/api/datasets/{snapshot.json()['id']}/download").status_code == 200
